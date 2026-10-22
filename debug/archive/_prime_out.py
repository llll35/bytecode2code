import math
from math import gcd
import os
try:
    from docx import Document
    from docx.shared import Inches
    docx_support = True
except ImportError:
    docx_support = False

class ReciprocalAnalyzer:

    def __init__(self, x):
        self.x = x
        self._validate_input()
        (self.a, self.b) = self._convert_to_fraction()
        self.non_repeating = self._get_non_repeating_part()
        self.repeating = self._get_repeating_part()

    def _validate_input(self):
        if self.x == 0:
            raise ValueError('输入不能为0')
        if isinstance(self.x, float):
            try:
                str_x = str(self.x).rstrip('0').rstrip('.') if '.' in str(self.x) else str(self.x)
                if 'e' in str_x:
                    raise ValueError('不支持非有限小数或高精度浮点数')
            except:
                raise ValueError('仅支持整数或有限小数输入')

    def _convert_to_fraction(self):
        if isinstance(self.x, float):
            s = str(self.x).split('.')
            integer_part = s[0] if s[0] else '0'
            decimal_part = s[1].rstrip('0') if len(s) > 1 else ''
            d = len(decimal_part)
            denominator = int(integer_part) * 10 ** d + (int(decimal_part) if d > 0 else 0)
            numerator = 10 ** d
        else:
            (numerator, denominator) = (1, int(self.x))
        common_divisor = gcd(numerator, denominator)
        a = numerator // common_divisor
        b = denominator // common_divisor
        return (a, b)

    def _prime_factorization_25(self, n):
        (m, n5) = (0, 0)
        while n % 2 == 0:
            m += 1
            n = n // 2
        while n % 5 == 0:
            n5 += 1
            n = n // 5
        return (m, n5, n)

    def _get_non_repeating_part(self):
        (m, n5, _) = self._prime_factorization_25(self.b)
        max_exp = max(m, n5)
        if max_exp == 0:
            return ''
        factor = 10 ** max_exp
        integer_part = self.a * factor // self.b
        non_rep_str = str(integer_part).zfill(max_exp)
        return non_rep_str[-max_exp:] if len(non_rep_str) > max_exp else non_rep_str

    def _get_repeating_part(self):
        (m, n5, k) = self._prime_factorization_25(self.b)
        if k == 1:
            return ''
        max_exp = max(m, n5)
        a_prime = self.a * 10 ** max_exp % self.b
        remainder_map = {}
        repeating_digits = []
        current_remainder = a_prime
        while current_remainder not in remainder_map:
            remainder_map[current_remainder] = len(repeating_digits)
            current_remainder *= 10
            digit = current_remainder // k
            repeating_digits.append(str(digit))
            current_remainder = current_remainder % k
            if current_remainder == 0:
                return ''
        cycle_start = remainder_map[current_remainder]
        return ''.join(repeating_digits[cycle_start:])

    def get_cyclic_string(self):
        if not self.repeating:
            return f'0.{self.non_repeating}' if self.non_repeating else '1.0'
        if len(self.repeating) == 1:
            cyclic = f'{self.repeating}˙'
        else:
            cyclic = f'˙{self.repeating}˙'
        return f'0.{self.non_repeating}{cyclic}' if self.non_repeating else f'0.{cyclic}'

    def get_results(self):
        return (self.non_repeating, self.repeating)

    def __str__(self):
        (non_rep, rep) = self.get_results()
        return f"1/{self.x} → 不循环: '{non_rep}', 循环: '{rep}'"

def write_txt(filename, results):
    try:
        with open(filename, 'w', encoding='utf-8') as f:
            f.write('\n'.join(results))
        return (True, f'结果已成功保存到：{os.path.abspath(filename)}')
    except Exception as e:
        return (False, f'TXT文件保存失败: {e}')

def write_docx(filename, results):
    if not docx_support:
        return (False, '缺少python-docx库，无法生成DOCX文件（可通过 pip install python-docx 安装）')
    try:
        doc = Document()
        doc.add_heading(results[0], level=1)
        doc.add_paragraph(results[1])
        for line in results[2:-1]:
            doc.add_paragraph(line)
        doc.add_paragraph(results[-1])
        doc.save(filename)
        return (True, f'结果已成功保存到：{os.path.abspath(filename)}')
    except Exception as e:
        return (False, f'DOCX文件保存失败: {e}')

def write_doc(filename, results):
    try:
        with open(filename, 'w', encoding='utf-8') as f:
            f.write('\n'.join(results))
        return (True, f'DOC文件（兼容TXT格式）已保存到：{os.path.abspath(filename)}（提示：DOC为旧格式，建议优先使用DOCX/TXT）')
    except Exception as e:
        return (False, f'DOC文件保存失败: {e}')

def main():
    while True:
        n = input('请输入数值范围(如1,10，用英文逗号分隔):')
        try:
            parts = n.split(',')
            if len(parts) != 2:
                raise ValueError('请输入两个数字，用逗号分隔')
            (min_n, max_n) = map(int, parts)
            if min_n <= 0 or max_n <= 0:
                raise ValueError('数值必须为正整数')
            if min_n > max_n:
                (min_n, max_n) = (max_n, min_n)
            break
        except ValueError as e:
            print(f'输入错误: {e}，请重新输入')
    results = []
    title = f'从 {min_n} 到 {max_n} 的倒数循环小数分析：'
    separator = '-' * len(title)
    results.append(title)
    results.append(separator)
    for i in range(min_n, max_n + 1):
        try:
            analyzer = ReciprocalAnalyzer(i)
            line = f'{analyzer}\t简写形式: {analyzer.get_cyclic_string()}'
            results.append(line)
        except ValueError as e:
            results.append(f'1/{i} 错误: {e}')
    results.append(separator)
    print('\n'.join(results))
    while True:
        save_choice = input('\n是否将结果保存到文件？(y/n): ').strip().lower()
        if save_choice in ('y', 'n'):
            break
        print("请输入 'y'（是）或 'n'（否）")
    if save_choice == 'y':
        while True:
            print('\n请选择文件扩展名：')
            print('1. TXT（兼容性好，无需额外依赖）')
            print('2. DOCX（Word新格式，需安装python-docx库）')
            print('3. DOC（Word旧格式，兼容TXT模式）')
            ext_choice = input('输入选项编号（1/2/3）: ').strip()
            if ext_choice in ('1', '2', '3'):
                ext_map = {'1': 'txt', '2': 'docx', '3': 'doc'}
                selected_ext = ext_map[ext_choice]
                break
            print('请输入正确的选项编号（1/2/3）')
        default_filename = f'倒数循环分析_{min_n}-{max_n}.{selected_ext}'
        filename = input(f'\n请输入保存文件名（默认: {default_filename}）: ').strip()
        if not filename:
            filename = default_filename
        else:
            filename = ''.join([c for c in filename if c not in '/\\:*?"<>|'])
            if not filename.lower().endswith(f'.{selected_ext}'):
                filename += f'.{selected_ext}'
        path = input('路径(文件夹，如D:/,默认当前目录):')
        if selected_ext == 'txt':
            (success, msg) = write_txt(path + filename, results)
        elif selected_ext == 'docx':
            (success, msg) = write_docx(path + filename, results)
        else:
            (success, msg) = write_doc(path + filename, results)
        print(msg)
if __name__ == '__main__':
    main()